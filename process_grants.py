"""
EMpower Grant Dashboard — Data Processing Script
================================================
Usage:
    python process_grants.py

Inputs (place in same folder as this script):
    GL_Report.xlsx           — Complete GL export from Sage Intacct
                               Sheet: 'General Ledger report (3)'
    Restricted_Grant_Master_Reviewed.xlsx  — Grant master briefing sheet
                               Sheet: 'Restricted_Grant_Master'

Output:
    grants_data.json         — Data file read by the web app (index.html)

Rules applied:
    - Income  : all 4xxx GL accounts, all dates, no grant-start filter
    - Expenses: all 7xxx GL accounts, from grant start date onwards only
    - Currency: Debit/Credit columns are USD (Sage auto-converts GBP/INR/HKD)
    - FY26    : 1 Jul 2025 – 30 Jun 2026
    - FY27    : 1 Jul 2026 – 30 Jun 2027
    - Functional Expense codes 30-XXX (excluding 30-000) identify each grant
"""

import json
import os
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, date

# ── Optional: auto-push to GitHub after generating JSON ─────────────────────
# Set GITHUB_PUSH = True and fill GITHUB_REPO if you have git configured
GITHUB_PUSH = False
GITHUB_REPO = ""   # e.g. "https://github.com/yourname/empower-dashboard.git"

# ── File paths ───────────────────────────────────────────────────────────────
SCRIPT_DIR  = os.path.dirname(os.path.abspath(__file__))
GL_FILE     = os.path.join(SCRIPT_DIR, "GL_Report.xlsx")
MASTER_FILE = os.path.join(SCRIPT_DIR, "Restricted_Grant_Master_Reviewed.xlsx")
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "grants_data.json")

GL_SHEET     = "General Ledger report (3)"
MASTER_SHEET = "Restricted_Grant_Master"

# ── FY boundaries ────────────────────────────────────────────────────────────
FY26_START = datetime(2025, 7, 1)
FY26_END   = datetime(2026, 6, 30)
FY27_START = datetime(2026, 7, 1)
FY27_END   = datetime(2027, 6, 30)

# ── GL account → budget category ─────────────────────────────────────────────
CATEGORY_MAP = {
    # Grantmaking
    "7455": "Grants", "7456": "Grants", "7457": "Grants", "7458": "Grants",
    "7459": "Grants", "7460": "Grants", "7461": "Grants", "7462": "Grants",
    "7463": "Grants",
    # Capacity Strengthening
    "7490": "Capacity Strengthening", "7491": "Capacity Strengthening",
    "7493": "Capacity Strengthening", "7495": "Capacity Strengthening",
    "7497": "Capacity Strengthening", "7498": "Capacity Strengthening",
    "7499": "Capacity Strengthening",
    # Communications
    "7492": "Communications",
    # Staffing
    "7005": "Staffing", "7006": "Staffing", "7007": "Staffing",
    "7008": "Staffing", "7105": "Staffing", "7120": "Staffing",
    "7135": "Staffing",
    # Travel
    "7357": "Travel", "7358": "Travel", "7359": "Travel", "7360": "Travel",
    "7361": "Travel", "7362": "Travel", "7363": "Travel", "7364": "Travel",
    "7365": "Travel", "7366": "Travel", "7367": "Travel", "7368": "Travel",
    "7374": "Travel", "7375": "Travel", "7376": "Travel",
    "7410": "Travel",  "7420": "Travel",
    "7568": "Travel", "7569": "Travel", "7570": "Travel", "7571": "Travel",
}

def get_category(gl_account):
    return CATEGORY_MAP.get(gl_account, "Overhead / Other")


# ── Date parser ───────────────────────────────────────────────────────────────
def parse_date(val):
    """Robust date parser — handles datetime objects, YYYY-MM-DD strings,
    MM/DD/YYYY strings, MM-DD-YYYY strings, 2-digit years."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day)
    s = str(val).strip()
    # ISO format: 2025-07-07 or 2025-07-07 00:00:00
    try:
        return datetime.strptime(s[:10], "%Y-%m-%d")
    except ValueError:
        pass
    # MM/DD/YYYY or MM-DD-YYYY or MM/DD/YY
    s = s.replace("-", "/")
    parts = s.split("/")
    if len(parts) == 3:
        try:
            m, d, y = int(parts[0]), int(parts[1]), int(parts[2])
            if y < 100:
                y += 2000
            return datetime(y, m, d)
        except ValueError:
            pass
    return None


def safe_float(val, default=0.0):
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", "").strip())
    except (ValueError, AttributeError):
        return default


def fmt_date(dt):
    """Return ISO date string or empty string."""
    if dt is None:
        return ""
    return dt.strftime("%Y-%m-%d")


# ── Load openpyxl ─────────────────────────────────────────────────────────────
try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    subprocess.check_call([sys.executable, "-m", "pip", "install",
                           "openpyxl", "--break-system-packages", "-q"])
    import openpyxl


# ═══════════════════════════════════════════════════════════════════════════════
# STEP 1 — Read grant master
# ═══════════════════════════════════════════════════════════════════════════════
def load_grant_master(path, sheet):
    print(f"Reading grant master: {os.path.basename(path)}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    if sheet not in wb.sheetnames:
        # Try first available sheet
        sheet = wb.sheetnames[0]
        print(f"  Warning: sheet not found, using '{sheet}'")

    ws = wb[sheet]
    grants = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue   # skip header
        vals = list(row)
        if not vals[2]:
            continue   # no code

        code = str(vals[2]).strip()
        if not code.startswith("30-"):
            continue

        grants[code] = {
            "code":           code,
            "short_name":     str(vals[0]).strip()  if vals[0] else "",
            "name":           str(vals[1]).strip()  if vals[1] else "",
            "fe_name":        str(vals[3]).strip()  if vals[3] else "",
            "entity":         str(vals[4]).strip()  if vals[4] else "",
            "entity_group":   str(vals[5]).strip()  if vals[5] else "",
            "start":          parse_date(vals[7]),
            "end":            parse_date(vals[8]),
            "budget_total":   safe_float(vals[9]),
            "budget_grants":  safe_float(vals[10]),
            "budget_cs":      safe_float(vals[11]),
            "budget_comms":   safe_float(vals[12]),
            "budget_staff":   safe_float(vals[13]),
            "budget_travel":  safe_float(vals[14]),
            "budget_oh":      safe_float(vals[15]),
            "lead":           str(vals[23]).strip() if len(vals) > 23 and vals[23] else "",
            "lead_email":     str(vals[24]).strip() if len(vals) > 24 and vals[24] else "",
            "orig_currency":  str(vals[27]).strip() if len(vals) > 27 and vals[27] else "USD",
            "purpose":        str(vals[31]).strip() if len(vals) > 31 and vals[31] else "",
            "rep_schedule":   str(vals[32]).strip() if len(vals) > 32 and vals[32] else "",
            "payment_sched":  str(vals[33]).strip() if len(vals) > 33 and vals[33] else "",
            "agreement_status": str(vals[29]).strip() if len(vals) > 29 and vals[29] else "",
        }

    print(f"  Loaded {len(grants)} grants")
    return grants


# ═══════════════════════════════════════════════════════════════════════════════
# STEP 2 — Read complete GL
# ═══════════════════════════════════════════════════════════════════════════════
def load_gl(path, sheet):
    print(f"Reading GL: {os.path.basename(path)}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    if sheet not in wb.sheetnames:
        sheet = wb.sheetnames[0]
        print(f"  Warning: sheet not found, using '{sheet}'")

    ws = wb[sheet]
    rows = []
    current_account = None
    current_account_name = ""
    import re
    pattern = re.compile(r"^(\d{4})\s*-\s*(.+?)(?:\s*\(Balance|$)")

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = list(row)
        if i < 6:
            continue

        # Section header row — identifies GL account
        if vals[0] is not None and vals[1] is None and vals[2] is None:
            s = str(vals[0])
            m = pattern.match(s)
            if m:
                current_account = m.group(1)
                current_account_name = m.group(2).strip()
            continue

        if len(vals) < 20:
            continue

        # Functional expense code — must be 30-XXX (not 30-000)
        fe = str(vals[10]).strip() if vals[10] else ""
        if not fe.startswith("30-") or fe == "30-000":
            continue

        try:
            debit  = float(vals[18]) if vals[18] is not None else 0.0
            credit = float(vals[19]) if vals[19] is not None else 0.0
        except (ValueError, TypeError):
            continue

        if debit == 0 and credit == 0:
            continue

        posted = parse_date(vals[0])
        if posted is None:
            continue

        ac_prefix = current_account[0] if current_account else ""
        ac_type = (
            "income"  if ac_prefix == "4" else
            "expense" if ac_prefix == "7" else
            "other"
        )

        rows.append({
            "fe":       fe,
            "ac":       current_account,
            "ac_name":  current_account_name,
            "posted":   posted,
            "debit":    debit,
            "credit":   credit,
            "memo":     str(vals[3])  if vals[3]  else "",
            "vendor":   str(vals[8])  if vals[8]  else "",
            "employee": str(vals[6])  if vals[6]  else "",
            "curr":     str(vals[16]).strip() if vals[16] else "",
            "type":     ac_type,
            "category": get_category(current_account),
        })

    print(f"  Loaded {len(rows)} GL rows (30-XXX, all currencies)")
    return rows


# ═══════════════════════════════════════════════════════════════════════════════
# STEP 3 — Build grant data objects
# ═══════════════════════════════════════════════════════════════════════════════
def build_grant_data(grants, gl_rows):
    print("Building grant data...")
    result = []

    for code, g in sorted(grants.items()):
        start = g["start"]
        if not start:
            print(f"  Skipping {code} — no start date in master")
            continue

        # All GL rows for this grant
        grant_rows = [r for r in gl_rows if r["fe"] == code]

        # ── Income: all dates, no start-date filter ──────────────────────────
        income_rows = [r for r in grant_rows if r["type"] == "income"]

        def inc_period(s, e):
            return sum(r["credit"] - r["debit"]
                       for r in income_rows if s <= r["posted"] <= e)

        fy26_income   = inc_period(FY26_START, FY26_END)
        fy27_income   = inc_period(FY27_START, FY27_END)
        total_income  = sum(r["credit"] - r["debit"] for r in income_rows)

        # Income line details for snapshot
        income_lines = [
            {
                "date":   fmt_date(r["posted"]),
                "memo":   r["memo"][:120],
                "vendor": r["vendor"][:80],
                "ac":     r["ac"],
                "amount": round(r["credit"] - r["debit"], 2),
            }
            for r in sorted(income_rows, key=lambda x: x["posted"])
        ]

        # ── Expenses: from grant start date only ─────────────────────────────
        expense_rows = [r for r in grant_rows
                        if r["type"] == "expense"
                        and r["posted"] >= start]

        def exp_period(s, e):
            return sum(r["debit"] - r["credit"]
                       for r in expense_rows if s <= r["posted"] <= e)

        fy26_expenses  = exp_period(FY26_START, FY26_END)
        fy27_expenses  = exp_period(FY27_START, FY27_END)
        total_expenses = sum(r["debit"] - r["credit"] for r in expense_rows)

        # ── Budget vs actuals by category ────────────────────────────────────
        budget_categories = {
            "Grants":                 g["budget_grants"],
            "Capacity Strengthening": g["budget_cs"],
            "Communications":         g["budget_comms"],
            "Staffing":               g["budget_staff"],
            "Travel":                 g["budget_travel"],
            "Overhead / Other":       g["budget_oh"],
        }

        actuals_by_cat = defaultdict(float)
        fy26_by_cat    = defaultdict(float)
        fy27_by_cat    = defaultdict(float)

        for r in expense_rows:
            cat = r["category"]
            net = r["debit"] - r["credit"]
            actuals_by_cat[cat] += net
            if FY26_START <= r["posted"] <= FY26_END:
                fy26_by_cat[cat] += net
            elif FY27_START <= r["posted"] <= FY27_END:
                fy27_by_cat[cat] += net

        budget_vs_actuals = []
        for cat, bud in budget_categories.items():
            act    = round(actuals_by_cat.get(cat, 0), 2)
            fy26_a = round(fy26_by_cat.get(cat, 0), 2)
            fy27_a = round(fy27_by_cat.get(cat, 0), 2)
            budget_vs_actuals.append({
                "category":      cat,
                "budget":        round(bud, 2),
                "fy26_actual":   fy26_a,
                "fy27_actual":   fy27_a,
                "total_actual":  act,
                "variance":      round(bud - act, 2),
                "pct_used":      round(act / bud * 100, 1) if bud > 0 else 0,
            })

        # ── GL line details (for detail tab) ─────────────────────────────────
        gl_lines = [
            {
                "date":     fmt_date(r["posted"]),
                "ac":       r["ac"],
                "ac_name":  r["ac_name"][:60],
                "category": r["category"],
                "memo":     r["memo"][:120],
                "vendor":   r["vendor"][:80],
                "employee": r["employee"][:60],
                "debit":    round(r["debit"], 2),
                "credit":   round(r["credit"], 2),
                "net":      round(r["debit"] - r["credit"], 2),
                "fy":       ("FY26" if FY26_START <= r["posted"] <= FY26_END
                              else "FY27" if FY27_START <= r["posted"] <= FY27_END
                              else "Other"),
            }
            for r in sorted(expense_rows, key=lambda x: x["posted"])
        ]

        # ── Days remaining ────────────────────────────────────────────────────
        today = datetime.today()
        end   = g["end"]
        days_remaining = (end - today).days if end else None
        grant_status = (
            "Active"   if end and today <= end else
            "Ended"    if end and today > end  else
            "Unknown"
        )

        budget_total = g["budget_total"]
        pct_used     = round(total_expenses / budget_total * 100, 1) if budget_total > 0 else 0

        result.append({
            # ── Identifiers ──────────────────────────────────────────────────
            "code":          code,
            "short_name":    g["short_name"],
            "name":          g["name"],
            "entity":        g["entity"],
            "entity_group":  g["entity_group"],

            # ── Grant details ─────────────────────────────────────────────────
            "start_date":        fmt_date(start),
            "end_date":          fmt_date(end),
            "days_remaining":    days_remaining,
            "grant_status":      grant_status,
            "lead":              g["lead"],
            "lead_email":        g["lead_email"],
            "orig_currency":     g["orig_currency"],
            "budget_total":      round(budget_total, 2),
            "purpose":           g["purpose"],
            "rep_schedule":      g["rep_schedule"],
            "payment_schedule":  g["payment_sched"],
            "agreement_status":  g["agreement_status"],

            # ── Income summary ────────────────────────────────────────────────
            "fy26_income":       round(fy26_income,  2),
            "fy27_income":       round(fy27_income,  2),
            "total_income":      round(total_income, 2),
            "income_lines":      income_lines,

            # ── Expense summary ───────────────────────────────────────────────
            "fy26_expenses":     round(fy26_expenses,  2),
            "fy27_expenses":     round(fy27_expenses,  2),
            "total_expenses":    round(total_expenses, 2),
            "balance":           round(budget_total - total_expenses, 2),
            "pct_used":          pct_used,

            # ── Budget vs actuals ─────────────────────────────────────────────
            "budget_categories": budget_vs_actuals,

            # ── GL line detail ────────────────────────────────────────────────
            "gl_lines":          gl_lines,
        })

        print(f"  {code:10}  {g['name'][:38]:38}  "
              f"Inc:{total_income:>12,.0f}  "
              f"Exp FY26:{fy26_expenses:>11,.0f}  "
              f"Exp FY27:{fy27_expenses:>11,.0f}  "
              f"Total Exp:{total_expenses:>12,.0f}  "
              f"{pct_used:>5.1f}%")

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# STEP 4 — Write JSON
# ═══════════════════════════════════════════════════════════════════════════════
def write_json(grant_data, path):
    output = {
        "generated":   datetime.today().strftime("%Y-%m-%d %H:%M"),
        "fy26_period": "01 Jul 2025 – 30 Jun 2026",
        "fy27_period": "01 Jul 2026 – 30 Jun 2027",
        "grants":      grant_data,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
    print(f"\nOutput written: {path}")
    size_kb = os.path.getsize(path) / 1024
    print(f"File size: {size_kb:.1f} KB")


# ═══════════════════════════════════════════════════════════════════════════════
# STEP 5 — Optional git push
# ═══════════════════════════════════════════════════════════════════════════════
def git_push(json_path):
    if not GITHUB_PUSH:
        return
    try:
        cmds = [
            ["git", "add", json_path],
            ["git", "commit", "-m",
             f"Update grant data {datetime.today().strftime('%Y-%m-%d %H:%M')}"],
            ["git", "push"],
        ]
        for cmd in cmds:
            result = subprocess.run(cmd, cwd=SCRIPT_DIR, capture_output=True, text=True)
            if result.returncode != 0:
                print(f"  Git warning: {result.stderr.strip()}")
            else:
                print(f"  {' '.join(cmd[:2])}: OK")
        print("Dashboard updated — web app will refresh within ~60 seconds.")
    except Exception as e:
        print(f"Git push failed: {e}")


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════
def main():
    print("=" * 70)
    print("EMpower Grant Dashboard — Data Processing Script")
    print("=" * 70)
    print()

    # Check files exist — try both master file names
    master_path = MASTER_FILE
    if not os.path.exists(master_path):
        alt = os.path.join(SCRIPT_DIR, "GL_Report.xlsx")
        if os.path.exists(alt):
            master_path = alt
            print(f"Using master from: {os.path.basename(alt)}")

    for f, label in [(GL_FILE, "GL Report"), (master_path, "Grant Master")]:
        if not os.path.exists(f):
            print(f"ERROR: {label} not found: {f}")
            print("Place GL_Report.xlsx and Restricted_Grant_Master_Reviewed.xlsx "
                  "in the same folder as this script.")
            sys.exit(1)

    # Check if master is inside GL_Report.xlsx (combined file)
    wb_check = openpyxl.load_workbook(GL_FILE, read_only=True)
    if MASTER_SHEET in wb_check.sheetnames:
        # Master is in the GL file — load both from same file
        grants   = load_grant_master(GL_FILE, MASTER_SHEET)
        gl_rows  = load_gl(GL_FILE, GL_SHEET)
    else:
        grants   = load_grant_master(master_path, MASTER_SHEET)
        gl_rows  = load_gl(GL_FILE, GL_SHEET)

    print()
    grant_data = build_grant_data(grants, gl_rows)
    print()
    write_json(grant_data, OUTPUT_FILE)
    git_push(OUTPUT_FILE)

    print()
    print("Done. Copy grants_data.json to your GitHub repo folder")
    print("alongside index.html and push — the web app will update automatically.")
    print("=" * 70)


if __name__ == "__main__":
    main()
